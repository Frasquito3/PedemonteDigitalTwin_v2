
from __future__ import annotations

from dataclasses import dataclass
from math import isclose
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from planner.core.config import ConfiguracionPlanificacion
from planner.core.schema import InstanciaTurno, PedidoInput, PlanTurno
from planner.domain.validator import validar_plan
from planner.routing.operations import (
    ResultadoOperacionEstimada,
    estimar_espera_cliente,
    simular_plan_operativo_estimado,
    tiempo_descarga_estimado_min,
)
from planner.routing.travel import MatrizViaje, tiempo_viaje_esperado_min


COLOR_AZUL = "17324D"
COLOR_CELESTE = "DCE6F0"
COLOR_BLANCO = "FFFFFF"
COLOR_FONDO = "F3F7FA"
COLOR_ALERTA = "B91C1C"


@dataclass(frozen=True)
class EntregaRutaEstimada:
    camion_id: int
    numero_viaje: int
    orden_visita: int
    pedido: PedidoInput
    inicio_carga_min: float
    salida_corralon_min: float
    llegada_cliente_min: float
    inicio_descarga_min: float
    fin_descarga_min: float
    regreso_corralon_min: float


def _hora(minuto: float) -> str:
    minutos = int(round(minuto))
    dias, minuto_dia = divmod(minutos, 1440)
    horas, resto = divmod(minuto_dia, 60)

    texto = f"{horas:02d}:{resto:02d}"
    if dias > 0:
        texto += f" (+{dias} d)"
    return texto


def _construir_entregas(
    instancia: InstanciaTurno,
    plan: PlanTurno,
    matriz: MatrizViaje,
    configuracion: ConfiguracionPlanificacion,
) -> tuple[list[EntregaRutaEstimada], ResultadoOperacionEstimada]:

    validacion = validar_plan(instancia, plan)
    if not validacion.valido:
        raise ValueError(
            "No se puede exportar un plan inválido: "
            + " | ".join(validacion.errores)
        )

    operacion = simular_plan_operativo_estimado(
        instancia,
        plan,
        matriz,
        configuracion,
    )

    cargas = {
        (carga.camion_id, carga.numero_viaje): carga
        for carga in operacion.cargas
    }

    pedidos = {
        pedido.pedido_id: pedido
        for pedido in instancia.pedidos
    }

    entregas: list[EntregaRutaEstimada] = []

    for camion in sorted(
        plan.camiones,
        key=lambda elemento: elemento.camion_id,
    ):
        ultimo_regreso = float(instancia.hora_inicio_turno_min)

        for viaje in camion.viajes:
            clave = (camion.camion_id, viaje.numero_viaje)

            if clave not in cargas:
                raise RuntimeError(
                    f"No existe carga estimada para {clave}."
                )

            carga = cargas[clave]

            minuto_actual = carga.minuto_fin
            nodo_actual = configuracion.id_nodo_corralon

            entregas_viaje = []

            for orden, pedido_id in enumerate(
                viaje.pedido_ids,
                start=1,
            ):
                pedido = pedidos[pedido_id]

                minuto_actual += tiempo_viaje_esperado_min(
                    matriz,
                    nodo_actual,
                    pedido_id,
                    minuto_actual,
                    configuracion,
                )

                llegada = minuto_actual

                espera = estimar_espera_cliente(
                    pedido,
                    llegada,
                    configuracion,
                )

                inicio_descarga = (
                    espera.minuto_inicio_descarga
                )

                fin_descarga = (
                    inicio_descarga
                    + tiempo_descarga_estimado_min(
                        pedido,
                        configuracion,
                    )
                )

                entregas_viaje.append(
                    (
                        orden,
                        pedido,
                        llegada,
                        inicio_descarga,
                        fin_descarga,
                    )
                )

                minuto_actual = fin_descarga
                nodo_actual = pedido_id

            regreso = (
                minuto_actual
                + tiempo_viaje_esperado_min(
                    matriz,
                    nodo_actual,
                    configuracion.id_nodo_corralon,
                    minuto_actual,
                    configuracion,
                )
            )

            for (
                orden,
                pedido,
                llegada,
                inicio_descarga,
                fin_descarga,
            ) in entregas_viaje:
                entregas.append(
                    EntregaRutaEstimada(
                        camion_id=camion.camion_id,
                        numero_viaje=viaje.numero_viaje,
                        orden_visita=orden,
                        pedido=pedido,
                        inicio_carga_min=carga.minuto_inicio,
                        salida_corralon_min=carga.minuto_fin,
                        llegada_cliente_min=llegada,
                        inicio_descarga_min=inicio_descarga,
                        fin_descarga_min=fin_descarga,
                        regreso_corralon_min=regreso,
                    )
                )

            ultimo_regreso = regreso

        fin_estimado = (
            operacion.finales_camiones_min[camion.camion_id]
        )

        if not isclose(
            ultimo_regreso,
            fin_estimado,
            abs_tol=1e-6,
            rel_tol=1e-9,
        ):
            raise RuntimeError(
                "La cronología exportada no coincide con "
                "el estimador operativo para el camión "
                f"{camion.camion_id}."
            )

    if len(entregas) != len(instancia.pedidos):
        raise RuntimeError(
            "La cantidad de entregas exportadas "
            "no coincide con los pedidos."
        )

    return entregas, operacion


def _estilo_titulo(hoja, ultima_columna: int) -> None:
    ultima = get_column_letter(ultima_columna)

    hoja.merge_cells(f"A1:{ultima}1")

    titulo = hoja["A1"]
    titulo.font = Font(
        name="Aptos",
        size=16,
        bold=True,
        color=COLOR_BLANCO,
    )
    titulo.fill = PatternFill(
        "solid",
        fgColor=COLOR_AZUL,
    )
    titulo.alignment = Alignment(
        vertical="center",
    )

    hoja.row_dimensions[1].height = 32


def _crear_resumen(
    libro: Workbook,
    instancia: InstanciaTurno,
    plan: PlanTurno,
    operacion: ResultadoOperacionEstimada,
    modo_solicitado: str,
) -> None:
    hoja = libro.active
    assert hoja is not None
    hoja.title = "Resumen"

    hoja["A1"] = "HOJA DE RUTA ESTIMADA"
    _estilo_titulo(hoja, 2)

    datos = [
        ("Instancia", instancia.instancia_id),
        ("Fecha de operación", instancia.fecha_operacion),
        ("Turno", instancia.turno.value),
        ("Método seleccionado", modo_solicitado),
        ("Algoritmo aplicado", plan.algoritmo.value),
        ("Camiones", instancia.cantidad_camiones),
        ("Pedidos / tareas", len(instancia.pedidos)),
        (
            "Pedidos originales",
            len({
                pedido.pedido_original_id or pedido.pedido_id
                for pedido in instancia.pedidos
            }),
        ),
        (
            "Unidades totales",
            sum(
                pedido.unidades_capacidad
                for pedido in instancia.pedidos
            ),
        ),
        ("Viajes planificados", operacion.viajes_totales),
        (
            "Fin estimado de la operación",
            _hora(max(operacion.finales_camiones_min)),
        ),
        ("Costo estimado del plan", plan.costo_estimado),
        (
            "Origen de los horarios",
            "Estimador operativo determinista de Python",
        ),
        (
            "Importante",
            "Los horarios no son resultados reales de AnyLogic.",
        ),
    ]

    for fila, (etiqueta, valor) in enumerate(datos, start=3):
        hoja.cell(fila, 1, etiqueta)
        hoja.cell(fila, 2, valor)

        hoja.cell(fila, 1).font = Font(
            bold=True,
            color=COLOR_AZUL,
        )

        if fila % 2 == 0:
            for columna in (1, 2):
                hoja.cell(fila, columna).fill = PatternFill(
                    "solid",
                    fgColor=COLOR_FONDO,
                )

    hoja.column_dimensions["A"].width = 33
    hoja.column_dimensions["B"].width = 66
    hoja.freeze_panes = "A3"


def _crear_hoja_camion(
    libro: Workbook,
    camion_id: int,
    entregas: list[EntregaRutaEstimada],
) -> None:
    hoja = libro.create_sheet(f"Camion {camion_id}")

    encabezados = [
        "Viaje",
        "Orden",
        "Pedido original",
        "Tarea",
        "Cliente",
        "Dirección",
        "Barrio",
        "Unidades",
        "Volcador",
        "Ventana horaria",
        "Inicio carga est.",
        "Salida corralón est.",
        "Llegada cliente est.",
        "Inicio descarga est.",
        "Fin descarga est.",
        "Regreso corralón est.",
        "Llegada prevista",
        "Observaciones",
    ]

    hoja["A1"] = f"HOJA DE RUTA - CAMIÓN {camion_id}"
    _estilo_titulo(hoja, len(encabezados))

    hoja["A2"] = (
        "Planificación estimada previa a la simulación"
    )

    for columna, encabezado in enumerate(encabezados, start=1):
        celda = hoja.cell(3, columna, encabezado)
        celda.font = Font(
            bold=True,
            color=COLOR_BLANCO,
        )
        celda.fill = PatternFill(
            "solid",
            fgColor=COLOR_AZUL,
        )
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    hoja.row_dimensions[3].height = 36

    for numero_fila, entrega in enumerate(entregas, start=4):
        pedido = entrega.pedido

        fuera_ventana = (
            entrega.llegada_cliente_min
            > pedido.hora_hasta_min
        )

        ventana = (
            f"{_hora(pedido.hora_desde_min)} - "
            f"{_hora(pedido.hora_hasta_min)}"
        )

        if not pedido.tiene_ventana_especifica:
            ventana = "Turno: " + ventana

        valores = [
            entrega.numero_viaje,
            entrega.orden_visita,
            pedido.pedido_original_id or pedido.pedido_id,
            pedido.pedido_id,
            pedido.cliente,
            pedido.direccion,
            pedido.barrio,
            pedido.unidades_capacidad,
            "Sí" if pedido.requiere_volcador else "No",
            ventana,
            _hora(entrega.inicio_carga_min),
            _hora(entrega.salida_corralon_min),
            _hora(entrega.llegada_cliente_min),
            _hora(entrega.inicio_descarga_min),
            _hora(entrega.fin_descarga_min),
            _hora(entrega.regreso_corralon_min),
            (
                "Fuera de ventana"
                if fuera_ventana
                else "En horario"
            ),
            pedido.observaciones,
        ]

        for columna, valor in enumerate(valores, start=1):
            celda = hoja.cell(numero_fila, columna, valor)

            celda.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            if numero_fila % 2 == 0:
                celda.fill = PatternFill(
                    "solid",
                    fgColor=COLOR_FONDO,
                )

        if fuera_ventana:
            hoja.cell(numero_fila, 17).font = Font(
                bold=True,
                color=COLOR_ALERTA,
            )

    anchos = [
        10, 10, 20, 19, 26, 36, 20, 12, 13,
        25, 22, 23, 23, 22, 21, 24, 21, 35,
    ]

    for columna, ancho in enumerate(anchos, start=1):
        letra = get_column_letter(columna)
        hoja.column_dimensions[letra].width = ancho

    hoja.freeze_panes = "E4"
    hoja.auto_filter.ref = (
        f"A3:R{max(3, hoja.max_row)}"
    )
    hoja.sheet_view.showGridLines = False


def exportar_hoja_ruta_excel(
    instancia: InstanciaTurno,
    plan: PlanTurno,
    matriz: MatrizViaje,
    configuracion: ConfiguracionPlanificacion,
    archivo: str | Path,
    *,
    modo_solicitado: str | None = None,
) -> Path:
    """
    Exporta la planificación estimada, no la ejecución real.

    No modifica la instancia, el plan ni los algoritmos.
    Utiliza la misma matriz y configuración que el
    estimador operativo existente.
    """
    ruta = Path(archivo).expanduser().resolve()

    if ruta.suffix.lower() != ".xlsx":
        raise ValueError("El archivo debe tener extensión .xlsx.")

    entregas, operacion = _construir_entregas(
        instancia,
        plan,
        matriz,
        configuracion,
    )

    libro = Workbook()

    _crear_resumen(
        libro,
        instancia,
        plan,
        operacion,
        modo_solicitado or plan.algoritmo.value,
    )

    for camion_id in range(instancia.cantidad_camiones):
        entregas_camion = [
            entrega
            for entrega in entregas
            if entrega.camion_id == camion_id
        ]

        _crear_hoja_camion(
            libro,
            camion_id,
            entregas_camion,
        )

    ruta.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    libro.save(ruta)
    libro.close()

    return ruta