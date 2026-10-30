"""
Generador simplificado de turnos y pedidos para evaluación de algoritmos.

Genera únicamente los atributos operativos necesarios para ejecutar los 5 algoritmos
y la simulación:
- pedido_id
- latitud, longitud
- unidades (tamaño entero del pedido)
- requiere_volcador ("SI" / "NO")
- tiene_ventana ("SI" / "NO"), hora_desde, hora_hasta
- turno ("MANANA" / "TARDE")
"""

from __future__ import annotations

import argparse
import csv
from dataclasses import asdict, dataclass
from datetime import datetime
import json
from math import ceil, cos, pi, sin
from pathlib import Path
import random
import sys
from typing import Any, Sequence

import pandas as pd

PYTHON_ROOT = Path(__file__).resolve().parents[2]
if str(PYTHON_ROOT) not in sys.path:
    sys.path.insert(0, str(PYTHON_ROOT))

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    OPENPYXL_DISPONIBLE = True
except ImportError:
    OPENPYXL_DISPONIBLE = False

# Rutas predeterminadas
DEFAULT_HISTORICO_CSV = PYTHON_ROOT / "data" / "processed" / "demanda_geografica.csv"
DEFAULT_GEOCODING_CSV = (
    PYTHON_ROOT / "data" / "processed" / "geocoding" / "direcciones_geolocalizadas.csv"
)
DEFAULT_PROBABILIDADES_CSV = (
    PYTHON_ROOT / "data" / "processed" / "estadistica" / "probabilidades_demanda.csv"
)
DEFAULT_CANT_PEDIDOS_CSV = PYTHON_ROOT / "data" / "raw" / "cant_pedidos_turnos.csv"
DEFAULT_OUTPUT_DIR = PYTHON_ROOT / "data" / "processed" / "pedidos_generados"

NOMBRE_METADATA_GENERACION = "metadata_generacion.json"

HORARIOS_TURNOS = {
    "MANANA": {
        "inicio_min": 7 * 60 + 30,  # 07:30 -> 450
        "fin_min": 12 * 60,         # 12:00 -> 720
        "inicio_str": "07:30",
        "fin_str": "12:00",
    },
    "TARDE": {
        "inicio_min": 14 * 60,      # 14:00 -> 840
        "fin_min": 17 * 60,         # 17:00 -> 1020
        "inicio_str": "14:00",
        "fin_str": "17:00",
    },
}

ENCABEZADOS_PEDIDOS = [
    "pedido_id",
    "latitud",
    "longitud",
    "unidades",
    "requiere_volcador",
    "tiene_ventana",
    "hora_desde",
    "hora_hasta",
    "turno",
]


@dataclass(frozen=True)
class UbicacionHistorica:
    """Coordenada histórica de referencia con su peso de demanda zonal."""

    latitud: float
    longitud: float
    peso_frecuencia: float


@dataclass
class PedidoGenerado:
    """Atributos estrictamente necesarios para planificar y simular un pedido."""

    pedido_id: str
    latitud: float
    longitud: float
    unidades: int
    requiere_volcador: str  # "SI" o "NO"
    tiene_ventana: str      # "SI" o "NO"
    hora_desde: str         # "HH:MM"
    hora_hasta: str         # "HH:MM"
    turno: str              # "MANANA" o "TARDE"

    def a_dict(self) -> dict[str, Any]:
        return asdict(self)


def generar_semilla_aleatoria() -> int:
    """Genera una semilla entera positiva aleatoria para reproducibilidad."""
    return random.SystemRandom().randint(100_000, 999_999_999)


def _minutos_a_hora_str(minutos: int) -> str:
    return f"{minutos // 60:02d}:{minutos % 60:02d}"


def estimar_tareas_pedidos(
    pedidos: Sequence[PedidoGenerado],
    capacidad_camion: int = 8,
) -> int:
    """Calcula la cantidad de tareas luego de dividir por la capacidad del camión."""
    if capacidad_camion <= 0:
        raise ValueError("capacidad_camion debe ser > 0.")
    return sum(ceil(p.unidades / capacidad_camion) for p in pedidos)


def cargar_cantidades_pedidos_turnos(
    ruta_csv: Path = DEFAULT_CANT_PEDIDOS_CSV,
) -> list[int]:
    """Lee `cant_pedidos_turnos.csv` y devuelve la lista de cantidades de pedidos por turno."""
    if not ruta_csv.exists():
        raise FileNotFoundError(f"No se encontró el archivo: {ruta_csv}")

    df = pd.read_csv(ruta_csv, encoding="utf-8-sig")
    cols_lower = {str(c).strip().lower(): c for c in df.columns}
    col_cant = cols_lower.get("cant_pedidos", df.columns[1] if len(df.columns) >= 2 else None)
    if col_cant is None:
        raise ValueError(f"El archivo {ruta_csv} debe tener la columna 'cant_pedidos'.")

    conteos = [
        int(round(float(v)))
        for v in df[col_cant]
        if pd.notna(v) and int(round(float(v))) > 0
    ]
    if not conteos:
        raise ValueError(f"No se encontraron conteos válidos en: {ruta_csv}")
    return conteos


def sortear_cantidad_pedidos_turno(
    conteos_historicos: Sequence[int],
    rng: random.Random,
    min_cota: int = 1,
    max_cota: int | None = None,
) -> int:
    """
    Elige al azar un conteo histórico de `cant_pedidos_turnos.csv` y aplica una
    variación discreta (-1, 0, +1) para no repetir rígidamente los mismos valores.
    """
    if not conteos_historicos:
        raise ValueError("conteos_historicos no puede estar vacío.")

    base = int(rng.choice(conteos_historicos))
    delta = rng.choices([-1, 0, 1], weights=[0.25, 0.50, 0.25], k=1)[0]
    piso = max(min_cota, min(conteos_historicos))
    techo = max(conteos_historicos)
    if max_cota is not None and max_cota > 0:
        techo = min(techo, max_cota)
        piso = min(piso, techo)

    return max(piso, min(techo, base + delta))


def cargar_probabilidades_demanda(
    ruta_probabilidades: Path = DEFAULT_PROBABILIDADES_CSV,
) -> tuple[dict[tuple[str, str], float], dict[str, float]]:
    """Lee `probabilidades_demanda.csv` y retorna pesos por `(ciudad, barrio)` y por `barrio`."""
    por_ciudad_barrio: dict[tuple[str, str], float] = {}
    por_barrio: dict[str, float] = {}

    if not ruta_probabilidades.exists():
        return por_ciudad_barrio, por_barrio

    df = pd.read_csv(ruta_probabilidades, encoding="utf-8-sig")
    cols = {str(c).strip().lower(): c for c in df.columns}
    col_ciudad = cols.get("ciudad", "")
    col_barrio = cols.get("barrio", "")
    col_prob = cols.get("probabilidad_demanda", "")

    if not col_barrio or not col_prob:
        return por_ciudad_barrio, por_barrio

    for _, fila in df.iterrows():
        ciudad = str(fila.get(col_ciudad, "")).strip().lower() if col_ciudad else ""
        barrio = str(fila.get(col_barrio, "")).strip().lower()
        if not barrio or pd.isna(fila.get(col_prob)):
            continue
        peso = float(fila[col_prob])
        if peso > 0:
            por_ciudad_barrio[(ciudad, barrio)] = (
                por_ciudad_barrio.get((ciudad, barrio), 0.0) + peso
            )
            por_barrio[barrio] = por_barrio.get(barrio, 0.0) + peso

    return por_ciudad_barrio, por_barrio


def cargar_y_analizar_historico(
    ruta_csv: Path = DEFAULT_GEOCODING_CSV,
    ruta_probabilidades: Path = DEFAULT_PROBABILIDADES_CSV,
) -> list[UbicacionHistorica]:
    """
    Carga coordenadas de referencia desde `direcciones_geolocalizadas.csv` (o `demanda_geografica.csv`)
    y las pondera con `probabilidades_demanda.csv`.
    """
    if not ruta_csv.exists():
        raise FileNotFoundError(f"No se encontró el archivo de ubicaciones en: {ruta_csv}")

    df = pd.read_csv(ruta_csv, encoding="utf-8-sig")
    cols = {str(c).strip().lower(): c for c in df.columns}

    col_lat = cols.get("latitud") or cols.get("lat")
    col_lon = cols.get("longitud") or cols.get("lon") or cols.get("lng")
    col_barrio = cols.get("barrio", "")
    col_ciudad = cols.get("ciudad") or cols.get("ciudad_detectada") or ""

    if not col_lat or not col_lon:
        raise ValueError(f"El CSV {ruta_csv} debe contener latitud y longitud.")

    por_cb, por_b = cargar_probabilidades_demanda(ruta_probabilidades)
    puntos_crudos: list[tuple[float, float, tuple[str, str], str]] = []
    conteo_cb: dict[tuple[str, str], int] = {}
    conteo_b: dict[str, int] = {}

    for _, fila in df.iterrows():
        try:
            lat = float(fila[col_lat])
            lon = float(fila[col_lon])
        except (ValueError, TypeError):
            continue

        barrio = str(fila.get(col_barrio, "")).strip().lower() if col_barrio else ""
        ciudad = str(fila.get(col_ciudad, "")).strip().lower() if col_ciudad else ""
        clave_cb = (ciudad, barrio)

        conteo_cb[clave_cb] = conteo_cb.get(clave_cb, 0) + 1
        conteo_b[barrio] = conteo_b.get(barrio, 0) + 1
        puntos_crudos.append((lat, lon, clave_cb, barrio))

    if not puntos_crudos:
        raise ValueError(f"No se pudieron extraer coordenadas válidas desde: {ruta_csv}")

    ubicaciones: list[UbicacionHistorica] = []
    for lat, lon, clave_cb, barrio in puntos_crudos:
        if clave_cb in por_cb:
            peso = max(1.0, (por_cb[clave_cb] / conteo_cb[clave_cb]) * 10_000.0)
        elif barrio in por_b:
            peso = max(1.0, (por_b[barrio] / conteo_b[barrio]) * 10_000.0)
        else:
            peso = 1.0
        ubicaciones.append(
            UbicacionHistorica(latitud=lat, longitud=lon, peso_frecuencia=peso)
        )

    return ubicaciones


def _aplicar_jitter_coordenadas(
    ubi: UbicacionHistorica,
    rng: random.Random,
) -> tuple[float, float]:
    """
    Desplaza aleatoriamente la coordenada de referencia entre ~40 m y ~145 m para
    no repetir nunca la latitud y longitud exacta de un pedido histórico.
    """
    radio_grados = rng.uniform(0.00035, 0.00130)
    angulo = rng.uniform(0.0, 2.0 * pi)

    lat_nueva = round(ubi.latitud + radio_grados * sin(angulo), 6)
    lon_nueva = round(ubi.longitud + (radio_grados * cos(angulo)) / 0.84, 6)

    if lat_nueva == round(ubi.latitud, 6) and lon_nueva == round(ubi.longitud, 6):
        lat_nueva = round(ubi.latitud + 0.00045, 6)
        lon_nueva = round(ubi.longitud - 0.00045, 6)

    return lat_nueva, lon_nueva


def _generar_ventana_horaria(
    inicio_turno_min: int,
    fin_turno_min: int,
    prob_ventana: float,
    rng: random.Random,
) -> tuple[str, str, str]:
    """
    Sortea la ventana horaria del pedido en una de las 4 opciones reales:
    1. Sin ventana: `tiene_ventana = "NO"`, de inicio a fin del turno.
    2. Antes de X: desde inicio del turno hasta X (< fin del turno).
    3. Después de X: desde X (> inicio del turno) hasta fin del turno.
    4. Entre X e Y: estrictamente dentro del turno (inicio < X < Y < fin).
    """
    inicio_str = _minutos_a_hora_str(inicio_turno_min)
    fin_str = _minutos_a_hora_str(fin_turno_min)

    if rng.random() >= prob_ventana:
        return "NO", inicio_str, fin_str

    tipo = rng.choices(
        ["ANTES_DE", "DESPUES_DE", "ENTRE_X_Y"],
        weights=[0.45, 0.40, 0.15],
        k=1,
    )[0]

    if tipo == "ANTES_DE":
        opciones_hasta = list(range(inicio_turno_min + 60, fin_turno_min, 15))
        hasta_min = rng.choice(opciones_hasta)
        return "SI", inicio_str, _minutos_a_hora_str(hasta_min)

    if tipo == "DESPUES_DE":
        opciones_desde = list(range(inicio_turno_min + 30, fin_turno_min - 45, 15))
        desde_min = rng.choice(opciones_desde)
        return "SI", _minutos_a_hora_str(desde_min), fin_str

    opciones_x = list(range(inicio_turno_min + 15, fin_turno_min - 75, 15))
    desde_min = rng.choice(opciones_x)
    opciones_y = [
        m for m in range(desde_min + 45, fin_turno_min, 15) if m - desde_min <= 120
    ]
    if not opciones_y:
        opciones_y = list(range(desde_min + 30, fin_turno_min, 15))
    hasta_min = rng.choice(opciones_y)
    return "SI", _minutos_a_hora_str(desde_min), _minutos_a_hora_str(hasta_min)


def _asignar_unidades_pedidos_turno(
    cantidad_pedidos: int,
    capacidad_camion: int,
    camiones_dia_promedio: float,
    rng: random.Random,
    max_tareas: int | None = None,
) -> list[int]:
    """
    Reparte unidades enteras (>= 1) entre los pedidos del turno de modo que por día
    salgan en promedio ~6 camiones llenos (48 unidades/día -> ~24 unidades/turno),
    pudiendo variar hasta ~9 camiones en días intensos.
    """
    if cantidad_pedidos <= 0:
        return []

    media_turno = max(1.0, camiones_dia_promedio / 2.0)
    camiones_turno = rng.triangular(
        max(1.0, media_turno * 0.55),
        max(media_turno + 1.0, media_turno * 1.60),
        media_turno * 0.90,
    )
    unidades_objetivo = max(
        cantidad_pedidos,
        int(round(camiones_turno * capacidad_camion)),
    )

    unidades = [1 for _ in range(cantidad_pedidos)]
    unidades_restantes = unidades_objetivo - cantidad_pedidos

    tareas_extra = (
        cantidad_pedidos
        if max_tareas is None
        else max(0, max_tareas - cantidad_pedidos)
    )
    topes: list[int] = []
    pesos: list[float] = []
    for _ in range(cantidad_pedidos):
        if tareas_extra > 0 and rng.random() < 0.10:
            topes.append(capacidad_camion * 2)
            tareas_extra -= 1
            pesos.append(4.0)
        else:
            topes.append(capacidad_camion)
            pesos.append(rng.choices([1.0, 2.0, 3.5], weights=[0.55, 0.30, 0.15], k=1)[0])

    for _ in range(unidades_restantes):
        candidatos = [i for i in range(cantidad_pedidos) if unidades[i] < topes[i]]
        if not candidatos:
            break
        elegido = rng.choices(candidatos, weights=[pesos[i] for i in candidatos], k=1)[0]
        unidades[elegido] += 1

    return unidades


def generar_pedidos_para_un_dia(
    ubicaciones: Sequence[UbicacionHistorica],
    cantidad_pedidos: int,
    turno: str = "MANANA",
    dia_numero: int = 1,
    prob_volcador: float | None = None,
    frecuencia_volcador_dias: float = 4.0,
    prob_ventana: float = 0.55,
    capacidad_camion: int = 8,
    camiones_dia_promedio: float = 6.0,
    rng: random.Random | None = None,
    max_tareas: int | None = None,
) -> list[PedidoGenerado]:
    """Genera los pedidos de un turno con únicamente los atributos requeridos por los algoritmos."""
    if rng is None:
        rng = random.Random()

    pesos = [u.peso_frecuencia for u in ubicaciones]
    seleccionadas = rng.choices(ubicaciones, weights=pesos, k=cantidad_pedidos)

    horario = HORARIOS_TURNOS.get(turno, HORARIOS_TURNOS["MANANA"])
    unidades_lista = _asignar_unidades_pedidos_turno(
        cantidad_pedidos=cantidad_pedidos,
        capacidad_camion=capacidad_camion,
        camiones_dia_promedio=camiones_dia_promedio,
        rng=rng,
        max_tareas=max_tareas,
    )

    if prob_volcador is not None:
        prob_volc = max(0.0, min(1.0, float(prob_volcador)))
    else:
        # 1 volcador cada 4 días = 1 cada 8 turnos -> esperanza 0.125 por turno
        prob_volc = min(0.50, (1.0 / max(1.0, 2.0 * frecuencia_volcador_dias)) / max(1, cantidad_pedidos))

    pedidos: list[PedidoGenerado] = []
    for idx, (ubi, unidades) in enumerate(zip(seleccionadas, unidades_lista), start=1):
        lat, lon = _aplicar_jitter_coordenadas(ubi, rng)
        requiere_volcador = "SI" if rng.random() < prob_volc else "NO"
        tiene_ventana, hora_desde, hora_hasta = _generar_ventana_horaria(
            inicio_turno_min=horario["inicio_min"],
            fin_turno_min=horario["fin_min"],
            prob_ventana=prob_ventana,
            rng=rng,
        )

        pedidos.append(
            PedidoGenerado(
                pedido_id=f"PED-D{dia_numero:02d}-{idx:02d}",
                latitud=lat,
                longitud=lon,
                unidades=int(unidades),
                requiere_volcador=requiere_volcador,
                tiene_ventana=tiene_ventana,
                hora_desde=hora_desde,
                hora_hasta=hora_hasta,
                turno=turno,
            )
        )

    return pedidos


def exportar_a_csv(pedidos: list[PedidoGenerado], ruta_archivo: Path) -> None:
    """Exporta los pedidos generados a CSV."""
    ruta_archivo.parent.mkdir(parents=True, exist_ok=True)
    with open(ruta_archivo, "w", encoding="utf-8-sig", newline="") as f:
        escritor = csv.DictWriter(f, fieldnames=ENCABEZADOS_PEDIDOS)
        escritor.writeheader()
        for p in pedidos:
            escritor.writerow(p.a_dict())


def exportar_a_excel(pedidos: list[PedidoGenerado], ruta_archivo: Path) -> None:
    """Exporta los pedidos generados a un archivo .xlsx compatible con `importar_pedidos_excel`."""
    if not OPENPYXL_DISPONIBLE:
        raise RuntimeError("openpyxl no está instalado.")

    ruta_archivo.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("Pedidos")
    else:
        ws.title = "Pedidos"

    ws.append(ENCABEZADOS_PEDIDOS)

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="17324D", end_color="17324D", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, len(ENCABEZADOS_PEDIDOS) + 1):
        celda = ws.cell(row=1, column=col_idx)
        celda.font = header_font
        celda.fill = header_fill
        celda.alignment = center_align

    for p in pedidos:
        ws.append([getattr(p, col) for col in ENCABEZADOS_PEDIDOS])

    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(max_len + 3, 12)

    wb.save(ruta_archivo)
    wb.close()


def parsear_rango_pedidos(valor: str) -> tuple[int, int] | None:
    """Parsea 'historico' -> None, '10' -> (10, 10) o '8-12' -> (8, 12)."""
    valor = valor.strip()
    if valor.lower() in ("historico", "hist", "auto", "csv"):
        return None
    if "-" in valor:
        partes = valor.split("-", 1)
        min_p, max_p = int(partes[0].strip()), int(partes[1].strip())
        if min_p <= 0 or max_p < min_p:
            raise ValueError(f"Rango de pedidos inválido: {valor}")
        return min_p, max_p
    p = int(valor)
    if p <= 0:
        raise ValueError(f"Cantidad de pedidos debe ser > 0: {valor}")
    return p, p


def ejecutar_generacion(
    dias_n: int = 1,
    rango_pedidos: tuple[int, int] | None = None,
    ruta_historico: Path = DEFAULT_GEOCODING_CSV,
    ruta_probabilidades: Path = DEFAULT_PROBABILIDADES_CSV,
    ruta_cant_pedidos: Path = DEFAULT_CANT_PEDIDOS_CSV,
    directorio_salida: Path = DEFAULT_OUTPUT_DIR,
    formato: str = "ambos",
    turno: str = "MANANA",
    prefijo: str = "pedidos_dia",
    crear_consolidado: bool = False,
    seed: int | None = None,
    usar_timestamp: bool = True,
    prob_volcador: float | None = None,
    frecuencia_volcador_dias: float = 4.0,
    prob_ventana: float = 0.55,
    capacidad_camion: int = 8,
    camiones_dia_promedio: float = 6.0,
    max_tareas: int | None = None,
) -> list[tuple[Path, list[PedidoGenerado]]]:
    """Genera N turnos y guarda los archivos y semillas en `seed_<seed_inicial>_<fecha>`."""
    if dias_n <= 0:
        raise ValueError(f"La cantidad de días (N) debe ser mayor a 0: {dias_n}")

    seed_inicial = int(seed) if seed is not None else generar_semilla_aleatoria()
    if seed_inicial <= 0:
        raise ValueError(f"La semilla inicial debe ser > 0: {seed_inicial}")

    rng = random.Random(seed_inicial)

    if not ruta_historico.exists() and DEFAULT_HISTORICO_CSV.exists():
        ruta_historico = DEFAULT_HISTORICO_CSV

    ubicaciones = cargar_y_analizar_historico(
        ruta_historico,
        ruta_probabilidades=ruta_probabilidades,
    )
    conteos_historicos = (
        cargar_cantidades_pedidos_turnos(ruta_cant_pedidos)
        if rango_pedidos is None
        else []
    )

    ahora = datetime.now()
    marca_tiempo = ahora.strftime("%Y%m%d_%H%M%S")
    carpeta_destino = (
        directorio_salida / f"seed_{seed_inicial}_{marca_tiempo}"
        if usar_timestamp
        else directorio_salida
    )
    carpeta_destino.mkdir(parents=True, exist_ok=True)

    resultados: list[tuple[Path, list[PedidoGenerado]]] = []
    todos_los_pedidos: list[PedidoGenerado] = []
    metadata_turnos: list[dict[str, Any]] = []

    for i in range(1, dias_n + 1):
        seed_escenario = rng.randint(100_000, 999_999_999)
        rng_dia = random.Random(seed_escenario)

        if rango_pedidos is not None:
            cant = rng_dia.randint(rango_pedidos[0], rango_pedidos[1])
        else:
            cant = sortear_cantidad_pedidos_turno(
                conteos_historicos=conteos_historicos,
                rng=rng_dia,
                max_cota=max_tareas,
            )

        turno_efectivo = (
            ("MANANA" if rng_dia.random() < 0.5 else "TARDE")
            if turno == "MIXTO"
            else turno
        )

        pedidos_dia = generar_pedidos_para_un_dia(
            ubicaciones=ubicaciones,
            cantidad_pedidos=cant,
            turno=turno_efectivo,
            dia_numero=i,
            prob_volcador=prob_volcador,
            frecuencia_volcador_dias=frecuencia_volcador_dias,
            prob_ventana=prob_ventana,
            capacidad_camion=capacidad_camion,
            camiones_dia_promedio=camiones_dia_promedio,
            rng=rng_dia,
            max_tareas=max_tareas,
        )

        nombre_base = f"{prefijo}_{i:02d}"
        archivo_csv_nombre: str | None = None
        archivo_xlsx_nombre: str | None = None
        ruta_principal_dia: Path | None = None

        if formato in ("csv", "ambos"):
            ruta_csv = carpeta_destino / f"{nombre_base}.csv"
            exportar_a_csv(pedidos_dia, ruta_csv)
            archivo_csv_nombre = ruta_csv.name
            ruta_principal_dia = ruta_csv

        if formato in ("xlsx", "ambos") and OPENPYXL_DISPONIBLE:
            ruta_xlsx = carpeta_destino / f"{nombre_base}.xlsx"
            exportar_a_excel(pedidos_dia, ruta_xlsx)
            archivo_xlsx_nombre = ruta_xlsx.name
            if ruta_principal_dia is None:
                ruta_principal_dia = ruta_xlsx

        if ruta_principal_dia is not None:
            resultados.append((ruta_principal_dia, pedidos_dia))

        tareas_estimadas = estimar_tareas_pedidos(
            pedidos_dia,
            capacidad_camion=capacidad_camion,
        )

        metadata_turnos.append(
            {
                "dia_numero": i,
                "instancia_id": f"TURNO-{i:02d}-SEED-{seed_escenario}",
                "seed_escenario": seed_escenario,
                "turno": turno_efectivo,
                "cantidad_pedidos": len(pedidos_dia),
                "tareas_estimadas": tareas_estimadas,
                "unidades_totales": sum(p.unidades for p in pedidos_dia),
                "archivo_csv": archivo_csv_nombre,
                "archivo_xlsx": archivo_xlsx_nombre,
                "pedidos": [p.a_dict() for p in pedidos_dia],
            }
        )
        todos_los_pedidos.extend(pedidos_dia)

    if crear_consolidado and dias_n > 1:
        if formato in ("csv", "ambos"):
            exportar_a_csv(todos_los_pedidos, carpeta_destino / f"{prefijo}_consolidado.csv")
        if formato in ("xlsx", "ambos") and OPENPYXL_DISPONIBLE:
            exportar_a_excel(todos_los_pedidos, carpeta_destino / f"{prefijo}_consolidado.xlsx")

    metadata_payload = {
        "version": 1,
        "seed_inicial": seed_inicial,
        "fecha_generacion": ahora.isoformat(timespec="seconds"),
        "marca_tiempo": marca_tiempo,
        "carpeta_salida": str(carpeta_destino),
        "parametros": {
            "dias_n": dias_n,
            "rango_pedidos": list(rango_pedidos) if rango_pedidos is not None else "historico",
            "turno_solicitado": turno,
            "formato": formato,
            "capacidad_camion": capacidad_camion,
            "camiones_dia_promedio": camiones_dia_promedio,
            "frecuencia_volcador_dias": frecuencia_volcador_dias,
            "prob_ventana": prob_ventana,
            "max_tareas": max_tareas,
        },
        "turnos": metadata_turnos,
    }

    ruta_metadata = carpeta_destino / NOMBRE_METADATA_GENERACION
    ruta_metadata.write_text(
        json.dumps(metadata_payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return resultados


def _parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Generador simplificado de turnos de pedidos para benchmarking."
    )
    p.add_argument("-n", "--dias", type=int, default=1, help="Cantidad de turnos a generar")
    p.add_argument(
        "-p",
        "--pedidos-por-dia",
        type=str,
        default="historico",
        help="Cantidad de pedidos por turno: 'historico', número fijo o rango (ej. 8-12)",
    )
    p.add_argument("--input-file", type=Path, default=DEFAULT_GEOCODING_CSV)
    p.add_argument("-o", "--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    p.add_argument("--sin-timestamp", action="store_true")
    p.add_argument("--formato", choices=["csv", "xlsx", "ambos"], default="ambos")
    p.add_argument("--turno", choices=["MANANA", "TARDE", "MIXTO"], default="MANANA")
    p.add_argument("--prefijo", type=str, default="pedidos_dia")
    p.add_argument("--consolidado", action="store_true")
    p.add_argument("--seed", type=int, default=None)
    p.add_argument("--max-tareas", type=int, default=None)
    return p


def main() -> int:
    args = _parser().parse_args()
    try:
        ejecutar_generacion(
            dias_n=args.dias,
            rango_pedidos=parsear_rango_pedidos(args.pedidos_por_dia),
            ruta_historico=args.input_file.resolve(),
            directorio_salida=args.output_dir.resolve(),
            formato=args.formato,
            turno=args.turno,
            prefijo=args.prefijo,
            crear_consolidado=args.consolidado,
            seed=args.seed,
            usar_timestamp=not args.sin_timestamp,
            max_tareas=args.max_tareas,
        )
        return 0
    except Exception as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
