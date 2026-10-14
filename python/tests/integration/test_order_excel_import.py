from __future__ import annotations
from datetime import time
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import Workbook
import pytest
from planner.integration.order_excel_import import ErrorImportacionExcel, formatear_minutos_hora, importar_pedidos_excel, parsear_hora_a_minutos
ENCABEZADOS = ['pedido_id', 'cliente', 'direccion', 'barrio', 'unidades', 'latitud', 'longitud', 'requiere_volcador', 'tiene_ventana', 'hora_desde', 'hora_hasta', 'observaciones']

def _crear_planilla(ruta: Path, filas: list[list[object]]) -> None:
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja.title = 'Pedidos'
    hoja.append(ENCABEZADOS)
    for fila in filas:
        hoja.append(fila)
    libro.save(ruta)
    libro.close()

def test_importa_hora_texto_y_hora_excel() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / 'pedidos.xlsx'
        _crear_planilla(ruta, [['P-1', 'Cliente 1', 'Dirección 1', 'Centro', 3, -32.85, -60.72, 'NO', 'SI', '07:45', time(9, 30), ''], ['P-2', 'Cliente 2', 'Dirección 2', 'Norte', 2, -32.84, -60.71, 'SI', 'NO', None, None, '']])
        resultado = importar_pedidos_excel(ruta, turno='MANANA')
        assert len(resultado.pedidos) == 2
        assert resultado.tareas_estimadas == 2
        assert resultado.pedidos[0].hora_desde_min == 465
        assert resultado.pedidos[0].hora_hasta_min == 570
        assert resultado.pedidos[1].hora_desde_min == 450
        assert resultado.pedidos[1].hora_hasta_min == 720

def test_split_estima_tareas_y_respeta_maximo() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / 'pedidos.xlsx'
        _crear_planilla(ruta, [['P-1', '', '', '', 17, -32.85, -60.72, 'SI', 'NO', None, None, '']])
        resultado = importar_pedidos_excel(ruta, turno='MANANA', capacidad_camion=8, max_tareas=3)
        assert resultado.tareas_estimadas == 3
        with pytest.raises(ErrorImportacionExcel):
            importar_pedidos_excel(ruta, turno='MANANA', capacidad_camion=8, max_tareas=2)

def test_rechaza_id_duplicado_y_ventana_fuera_de_turno() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / 'pedidos.xlsx'
        _crear_planilla(ruta, [['P-1', '', '', '', 2, -32.85, -60.72, 'NO', 'SI', '07:00', '08:00', ''], ['P-1', '', '', '', 2, -32.84, -60.71, 'NO', 'NO', None, None, '']])
        with pytest.raises(ErrorImportacionExcel) as captura:
            importar_pedidos_excel(ruta, turno='MANANA')
        mensajes = [error.mensaje for error in captura.value.errores]
        assert any(('dentro del turno' in mensaje for mensaje in mensajes))
        assert any(('ID duplicado' in mensaje for mensaje in mensajes))


# La columna nueva es optativa: un archivo sin ella mantiene su semántica.
ENCABEZADOS_MIXTOS = [*ENCABEZADOS, "unidades_volcador"]


def _crear_planilla_mixta(ruta: Path, filas: list[list[object]], *, encabezado: str = "unidades_volcador") -> None:
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja.title = "Pedidos"
    hoja.append([*ENCABEZADOS, encabezado])
    for fila in filas:
        hoja.append(fila)
    libro.save(ruta)
    libro.close()


def _fila_mixta(
    *,
    unidades: int = 9,
    requiere: str = "SI",
    cantidad: object = 4,
) -> list[object]:
    return [
        "MIX-01", "Cliente real", "Calle 123", "Centro",
        unidades, -32.85, -60.72, requiere, "NO", None, None,
        "Entrega en obra", cantidad,
    ]


def test_mixto_9_con_4_genera_2_tareas() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "mix.xlsx"
        _crear_planilla_mixta(ruta, [_fila_mixta()])
        resultado = importar_pedidos_excel(ruta, turno="MANANA", capacidad_camion=8)
        assert resultado.tareas_estimadas == 2
        assert resultado.pedidos[0].unidades_volcador == 4
        assert resultado.pedidos[0].requiere_volcador
        assert resultado.pedidos[0].cliente == "Cliente real"


def test_mixto_8_con_3_tambien_genera_2_tareas() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "mix.xlsx"
        _crear_planilla_mixta(ruta, [_fila_mixta(unidades=8, cantidad=3)])
        resultado = importar_pedidos_excel(ruta, turno="MANANA")
        assert resultado.tareas_estimadas == 2


def test_mixto_20_con_11_genera_4_tareas() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "mix.xlsx"
        _crear_planilla_mixta(ruta, [_fila_mixta(unidades=20, cantidad=11)])
        resultado = importar_pedidos_excel(ruta, turno="MANANA")
        assert resultado.tareas_estimadas == 4
        with pytest.raises(ErrorImportacionExcel, match="supera el máximo"):
            importar_pedidos_excel(ruta, turno="MANANA", max_tareas=3)


def test_legacy_si_significa_todas_las_unidades_volcador() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "legacy.xlsx"
        _crear_planilla(ruta, [_fila_mixta(unidades=9)[:-1]])
        resultado = importar_pedidos_excel(ruta, turno="MANANA")
        assert resultado.tareas_estimadas == 2
        assert resultado.pedidos[0].unidades_volcador == 9


def test_legacy_no_significa_cero_volcador() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "legacy.xlsx"
        _crear_planilla(ruta, [_fila_mixta(unidades=7, requiere="NO")[:-1]])
        resultado = importar_pedidos_excel(ruta, turno="MANANA")
        assert resultado.pedidos[0].unidades_volcador == 0
        assert resultado.tareas_estimadas == 1


def test_nuevo_no_permite_cantidad_blanca_si_requiere_volcador() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "mix.xlsx"
        _crear_planilla_mixta(ruta, [_fila_mixta(cantidad=None)])
        with pytest.raises(ErrorImportacionExcel) as ex:
            importar_pedidos_excel(ruta, turno="MANANA")
        assert any(e.campo == "unidades_volcador" and e.fila == 2 for e in ex.value.errores)


def test_nuevo_no_permite_mayor_que_total() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "mix.xlsx"
        _crear_planilla_mixta(ruta, [_fila_mixta(cantidad=10)])
        with pytest.raises(ErrorImportacionExcel, match="superar"):
            importar_pedidos_excel(ruta, turno="MANANA")


@pytest.mark.parametrize("cantidad", [-1, 1.5, "dos", True])
def test_nuevo_rechaza_cantidad_no_entera_no_negativa(cantidad: object) -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "mix.xlsx"
        _crear_planilla_mixta(ruta, [_fila_mixta(cantidad=cantidad)])
        with pytest.raises(ErrorImportacionExcel, match="entero mayor o igual"):
            importar_pedidos_excel(ruta, turno="MANANA")


@pytest.mark.parametrize("requiere,cantidad", [("NO", 4), ("SI", 0)])
def test_nuevo_rechaza_contradiccion(requiere: str, cantidad: int) -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "mix.xlsx"
        _crear_planilla_mixta(ruta, [_fila_mixta(requiere=requiere, cantidad=cantidad)])
        with pytest.raises(ErrorImportacionExcel, match="requiere_volcador es SI"):
            importar_pedidos_excel(ruta, turno="MANANA")


def test_nuevo_admite_vacio_si_no_requiere_y_asigna_cero() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "mix.xlsx"
        _crear_planilla_mixta(ruta, [_fila_mixta(requiere="NO", cantidad=None)])
        resultado = importar_pedidos_excel(ruta, turno="MANANA")
        assert resultado.pedidos[0].unidades_volcador == 0
        assert resultado.tareas_estimadas == 2


def test_alias_columna_nueva_con_espacios() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "mix.xlsx"
        _crear_planilla_mixta(ruta, [_fila_mixta()], encabezado="Unidades con Volcador")
        resultado = importar_pedidos_excel(ruta, turno="MANANA")
        assert resultado.pedidos[0].unidades_volcador == 4


def test_roundtrip_formato_canonico_web_con_mixto():
    from tempfile import TemporaryDirectory
    from openpyxl import Workbook
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "pedidos_web.xlsx"
        libro = Workbook()
        hoja = libro.active
        assert hoja is not None
        hoja.title = "Pedidos"
        hoja.append([
            "pedido_id", "cliente", "direccion", "barrio", "unidades",
            "latitud", "longitud", "requiere_volcador",
            "unidades_volcador", "tiene_ventana", "hora_desde",
            "hora_hasta", "observaciones", "turno",
        ])
        hoja.append(["MIX", "Ana", "Calle 123", "Centro", 9, -32.85, -60.72,
                     "SI", 4, "NO", "", "", "Prueba", "MANANA"])
        hoja.append(["NOR", "Beto", "Calle 456", "Sur", 2, -32.84, -60.71,
                     "NO", 0, "SI", "08:00", "09:00", "Observacion", "MANANA"])
        libro.save(ruta)
        libro.close()
        resultado = importar_pedidos_excel(ruta, turno="MANANA")
        assert len(resultado.pedidos) == 2
        assert resultado.tareas_estimadas == 3
        assert [(p.unidades, p.unidades_volcador) for p in resultado.pedidos] == [(9, 4), (2, 0)]
        assert resultado.pedidos[1].tiene_ventana is True

def test_compatibilidad_exportacion_web_antigua_sin_tiene_ventana():
    from tempfile import TemporaryDirectory
    from openpyxl import Workbook
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "web_antiguo.xlsx"
        libro = Workbook()
        hoja = libro.active
        assert hoja is not None
        hoja.title = "Pedidos"
        hoja.append(["pedido_id", "unidades_capacidad", "latitud", "longitud",
                     "requiere_volcador", "hora_desde", "hora_hasta", "turno"])
        hoja.append(["ANT", 4, -32.85, -60.72, 1, "08:00", "09:00", "MANANA"])
        libro.save(ruta)
        libro.close()
        res = importar_pedidos_excel(ruta, turno="MANANA")
        assert res.pedidos[0].unidades_volcador == 4
        assert res.pedidos[0].tiene_ventana is True

def test_turno_exportado_distinto_se_rechaza():
    from tempfile import TemporaryDirectory
    from openpyxl import Workbook
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "turno.xlsx"
        libro = Workbook()
        hoja = libro.active
        assert hoja is not None
        hoja.title = "Pedidos"
        hoja.append(["pedido_id", "unidades", "latitud", "longitud",
                     "requiere_volcador", "tiene_ventana", "hora_desde", "hora_hasta", "turno"])
        hoja.append(["TAR", 4, -32.85, -60.72, "NO", "NO", "", "", "TARDE"])
        libro.save(ruta)
        libro.close()
        with pytest.raises(ErrorImportacionExcel) as error:
            importar_pedidos_excel(ruta, turno="MANANA")
        assert any(e.campo == "turno" for e in error.value.errores)
