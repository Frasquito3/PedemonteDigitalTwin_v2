from __future__ import annotations
from pathlib import Path
from tempfile import TemporaryDirectory
from urllib.parse import unquote
from openpyxl import Workbook
from planner.integration.order_import_bridge import importar_pedidos_excel_pypeline

def _guardar_planilla(ruta: Path, *, id_pedido: str='OP-01', hora_desde: str='07:30') -> None:
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Pedidos'
    hoja.append(['pedido_id', 'cliente', 'direccion', 'barrio', 'unidades', 'latitud', 'longitud', 'requiere_volcador', 'tiene_ventana', 'hora_desde', 'hora_hasta', 'observaciones'])
    hoja.append([id_pedido, 'Cliente con espacio', 'Ruta 11 | km 3', 'Norte', 4, -32.85, -60.72, 'SI', 'SI', hora_desde, '09:00', 'Prueba ñ'])
    libro.save(ruta)
    libro.close()

def test_protocolo_ok_es_parseable() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / 'pedidos.xlsx'
        _guardar_planilla(ruta)
        respuesta = importar_pedidos_excel_pypeline(str(ruta), 'MANANA')
        lineas = respuesta.splitlines()
        cabecera = lineas[0].split('|')
        pedido = lineas[1].split('|')
        assert cabecera[:5] == ['PEDXLSX1', 'OK', '1', '1', '0']
        assert pedido[0] == 'P'
        assert unquote(pedido[2]) == 'OP-01'
        assert unquote(pedido[4]) == 'Ruta 11 | km 3'
        assert pedido[9:13] == ['1', '1', '450', '540']
        assert unquote(pedido[13]) == 'Prueba ñ'

def test_protocolo_error_no_lanza_excepcion() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / 'pedidos.xlsx'
        _guardar_planilla(ruta, hora_desde='06:00')
        respuesta = importar_pedidos_excel_pypeline(str(ruta), 'MANANA')
        lineas = respuesta.splitlines()
        assert lineas[0].startswith('PEDXLSX1|ERROR|')
        assert lineas[1].startswith('E|')


def test_protocolo_v1_rechaza_mixto_hasta_actualizar_anylogic() -> None:
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "mix.xlsx"
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Pedidos"
        hoja.append([
            "pedido_id", "cliente", "direccion", "barrio", "unidades",
            "latitud", "longitud", "requiere_volcador", "tiene_ventana",
            "hora_desde", "hora_hasta", "observaciones", "unidades_volcador",
        ])
        hoja.append([
            "MIX", "Cliente", "Ruta 11", "Norte", 9,
            -32.85, -60.72, "SI", "NO", None, None, "", 4,
        ])
        libro.save(ruta)
        libro.close()
        respuesta = importar_pedidos_excel_pypeline(str(ruta), "MANANA")
        lineas = respuesta.splitlines()
        assert lineas[0] == "PEDXLSX1|ERROR|1"
        assert lineas[1].startswith("E|2|unidades_volcador|")
        assert "AnyLogic" in unquote(lineas[1])


def test_protocolo_v2_conserva_unidades_volcador():
    from planner.integration.order_import_bridge import importar_pedidos_excel_pypeline_v2
    with TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "mix_v2.xlsx"
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Pedidos"
        hoja.append(["pedido_id", "unidades", "latitud", "longitud", "requiere_volcador",
                     "unidades_volcador", "tiene_ventana", "hora_desde", "hora_hasta"])
        hoja.append(["MIX", 9, -32.85, -60.72, "SI", 4, "NO", "", ""])
        libro.save(ruta)
        libro.close()
        respuesta = importar_pedidos_excel_pypeline_v2(str(ruta), "MANANA")
        cab, pedido = respuesta.splitlines()
        assert cab.startswith("PEDXLSX2|OK|1|2|")
        partes = pedido.split("|")
        assert len(partes) == 15
        assert partes[9] == "1"
        assert partes[14] == "4"
